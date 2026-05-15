from __future__ import annotations
from datetime import datetime
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from app.models.session import ProFormaSession
from app.excel.cell_refs import (
    BLDG_ROW, BLDG_VAL_COL, TENANT_HEADER_ROW,
    TENANT_DATA_START_ROW, TENANT_COL,
)

_LABEL_FILL = PatternFill("solid", fgColor="D9D9D9")
_HDR_FILL = PatternFill("solid", fgColor="4F81BD")
_BOLD = Font(bold=True)
_BOLD_WHITE = Font(bold=True, color="FFFFFF")

# (label, key) in row order — index+1 matches BLDG_ROW row number
_BLDG_PARAMS = [
    ("Building Name",    "building_name"),
    ("Start Year",       "start_year"),
    ("Start Month",      "start_month"),
    ("Years to Project", "years"),
    ("Total SF",         "total_sqft"),
    ("Occupied SF",      "occupied_sqft"),
    ("OpEx / SF",        "opex_psf"),
    ("OpEx Growth Rate", "opex_growth"),
    ("Market Avg Rate",  "market_avg_rate"),
    ("Market Growth %",  "market_growth"),
    ("Cap Rate",         "cap_rate"),
    ("Cap Delta",        "cap_delta"),
    ("Weighted Avg Rate","weighted_avg_rate"),
]

_BLDG_VALUES = {
    "building_name": lambda s: s.building_name,
    "start_year":    lambda s: s.start_year,
    "start_month":   lambda s: s.start_month,
    "years":         lambda s: s.years,
    "total_sqft":    lambda s: s.total_sqft,
    "occupied_sqft": lambda s: s.occupied_sqft,
    "opex_psf":      lambda s: s.opex_psf,
    "opex_growth":   lambda s: 0.025,
    "market_avg_rate": lambda s: s.market_avg_rate,
    "market_growth": lambda s: s.market_growth_pct,
    "cap_rate":      lambda s: s.cap_rate,
    "cap_delta":     lambda s: s.cap_delta,
}


def write_inputs_sheet(ws, session: ProFormaSession) -> None:
    ws.title = "Inputs"
    num_t = len(session.tenants)
    first_t = TENANT_DATA_START_ROW
    last_t = TENANT_DATA_START_ROW + num_t - 1
    sf_col = get_column_letter(TENANT_COL["sqft"])
    rt_col = get_column_letter(TENANT_COL["rate_psf"])

    for label, key in _BLDG_PARAMS:
        row = BLDG_ROW[key]
        lbl_cell = ws.cell(row=row, column=1, value=label)
        lbl_cell.font = _BOLD
        lbl_cell.fill = _LABEL_FILL
        if key == "weighted_avg_rate":
            if num_t > 0:
                sf_r = f"{sf_col}{first_t}:{sf_col}{last_t}"
                rt_r = f"{rt_col}{first_t}:{rt_col}{last_t}"
                ws.cell(row=row, column=BLDG_VAL_COL,
                        value=f"=SUMPRODUCT({sf_r},{rt_r})/SUM({sf_r})")
        else:
            ws.cell(row=row, column=BLDG_VAL_COL, value=_BLDG_VALUES[key](session))

    for col_name, col_idx in TENANT_COL.items():
        c = ws.cell(row=TENANT_HEADER_ROW, column=col_idx,
                    value=col_name.replace("_", " ").title())
        c.font = _BOLD_WHITE; c.fill = _HDR_FILL
        c.alignment = Alignment(horizontal="center")

    for i, tenant in enumerate(session.tenants):
        row = TENANT_DATA_START_ROW + i
        values = {
            "name": tenant.name, "suite": tenant.suite,
            "sqft": tenant.sqft, "rate_psf": tenant.rate_psf,
            "year1_override": tenant.year1_override if tenant.year1_override is not None else "",
            "lease_exp": datetime.strptime(tenant.lease_exp, "%m-%d-%Y") if tenant.lease_exp and tenant.lease_exp.strip() else "",
            "projection_type": tenant.projection_type,
            "growth_rate": tenant.growth_rate, "flat_increase": tenant.flat_increase,
            "pct_increase": tenant.pct_increase, "renewed": tenant.renewed,
            "renewal_start": datetime.strptime(tenant.renewal_start, "%m-%d-%Y") if tenant.renewal_start and tenant.renewal_start.strip() else "",
            "renewal_term_years": tenant.renewal_term_years,
            "renewal_projection_type": tenant.renewal_projection_type,
            "renewal_growth_rate": tenant.renewal_growth_rate,
            "renewal_flat_increase": tenant.renewal_flat_increase,
            "renewal_pct_increase": tenant.renewal_pct_increase,
        }
        for col_name, col_idx in TENANT_COL.items():
            cell = ws.cell(row=row, column=col_idx, value=values[col_name])
            if col_name in ("lease_exp", "renewal_start") and values[col_name]:
                cell.number_format = "MM-DD-YYYY"

    for col in ws.columns:
        max_len = max(
            (len(str(c.value)) if c.value and not str(c.value).startswith("=") else 0)
            for c in col
        )
        ws.column_dimensions[get_column_letter(col[0].column)].width = max_len + 4
